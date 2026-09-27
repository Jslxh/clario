import logging
import docx
from docx.opc.exceptions import PackageNotFoundError

from app.services.parsers.base import BaseParser
from app.schemas.parser import ParsedDocument, ParsedPage
from app.core.exceptions import CorruptFileError

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    """Parser for extracting normalized paragraph and table text from DOCX documents."""

    def parse(self, file_path: str, document_id: str, filename: str) -> ParsedDocument:
        try:
            doc = docx.Document(file_path)
        except (PackageNotFoundError, Exception) as err:
            logger.error(f"Failed to read DOCX file '{filename}': {err}")
            raise CorruptFileError(f"File '{filename}' is corrupt or not a valid DOCX document.", document_type="docx") from err

        text_blocks = []
        sections = []

        # Process paragraphs preserving order and headings
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            if style_name and "Heading" in style_name:
                sections.append(text)
                text_blocks.append(f"## {text}")
            else:
                text_blocks.append(text)

        # Process tables conservatively
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_lines.append(f"| {row_text} |")
            if table_lines:
                text_blocks.extend(table_lines)

        full_text = "\n\n".join(text_blocks)
        has_usable_text = bool(full_text.strip())

        pages = [
            ParsedPage(
                page_number=1,
                text=full_text,
                sections=sections,
            )
        ]

        return ParsedDocument(
            document_id=document_id,
            filename=filename,
            document_type="docx",
            has_usable_text=has_usable_text,
            total_pages=1,
            pages=pages,
            metadata={
                "extracted_headings_count": len(sections),
                "extracted_tables_count": len(doc.tables),
            },
        )
